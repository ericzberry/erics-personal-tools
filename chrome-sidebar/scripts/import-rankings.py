"""Read the source workbook without modifying it; preserve each source's meaning."""
import argparse, hashlib, json
from pathlib import Path
import openpyxl
p=argparse.ArgumentParser()
p.add_argument('workbook',type=Path)
p.add_argument('--output',type=Path,default=Path('config/rankings-2026.json'))
a=p.parse_args()
w=openpyxl.load_workbook(a.workbook,data_only=True)
s=w['Combined Ranks']; players=[]; tier=1; previous_bottom=False
for row in s.iter_rows(min_row=2):
    if row[2].value is None: continue
    rank,name,pos,team,adp=[c.value for c in row[1:6]]
    if not isinstance(rank,int) or not name or not pos or not team: raise ValueError(f'Invalid ranking at row {row[0].row}')
    top=any(c.border.top.style in ('medium','thick') for c in row[2:6])
    if players and (top or previous_bottom): tier+=1
    previous_bottom=any(c.border.bottom.style in ('medium','thick') for c in row[2:6])
    players.append({'tier':tier,'rank':rank,'name':name,'position':pos,'nflTeam':team,'adp':adp,'highlighted':row[2].fill.fgColor.type=='rgb' and row[2].fill.fgColor.rgb=='FFFFFF00'})
if sorted(x['rank'] for x in players)!=list(range(1,len(players)+1)): raise ValueError('Ranks must be unique and sequential')
s=w['Ranks By Postition']; comparisons=[]
for start in [2,9,16,23]:
    for row in s.iter_rows(min_row=2):
        vals=[c.value for c in row[start-1:start+5]]
        if vals[3] is None: continue
        mine,boone,ff4,name,pos,team=vals
        comparisons.append({'name':name,'position':pos,'nflTeam':team,'me':mine,'boone':boone,'fourForFour':ff4})
a.output.parent.mkdir(parents=True,exist_ok=True)
a.output.write_text(json.dumps({'schemaVersion':2,'seasonId':2026,'scoring':'half-ppr','source':a.workbook.name,'sha256':hashlib.sha256(a.workbook.read_bytes()).hexdigest(),'priority':'Combined Ranks tiers (dark horizontal borders), then overall order, are authoritative. ADP is market context. Highlight meaning is unspecified.','players':players,'positionalComparisons':comparisons},indent=2)+'\n')
print(f'Imported {len(players)} overall entries and {len(comparisons)} positional comparisons.')
